from __future__ import annotations

from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path

from fastapi import Depends, FastAPI, HTTPException, Query, Response, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, RedirectResponse
import httpx
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.database import SessionLocal, create_schema, get_db
from app.diagnosis import build_diagnoses
from app.models import (
    Account,
    AccountMetricSnapshot,
    AccountStatus,
    AuditLog,
    Diagnosis,
    ExportJob,
    JobStatus,
    MetricRequirement,
    OAuthState,
    Role,
    SyncJob,
    User,
    Video,
    VideoMetricSnapshot,
    Workspace,
)
from app.douyin import authorization_url, encrypt, exchange_code, expiration, new_state, state_hash
from app.schemas import AccountCreate, AccountDataImport, AccountPatch, AccountSnapshotImport, ExportRequest, LoginRequest, VideoImport
from app.security import create_access_token, current_user, hash_password, require_role, verify_password

REQUIRED_METRICS = {
    "play_count": "video",
    "like_count": "video",
    "comment_count": "video",
    "forward_count": "video",
    "follower_delta": "account",
    "completion_rate": "video",
    "video_follow_conversion": "video",
}


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


async def get_workspace(session: AsyncSession) -> Workspace:
    workspace = await session.scalar(select(Workspace).limit(1))
    if not workspace:
        raise HTTPException(status_code=500, detail="工作区尚未初始化")
    return workspace


async def get_account(session: AsyncSession, account_id: str) -> Account:
    account = await session.scalar(select(Account).where(Account.id == account_id))
    if not account:
        raise HTTPException(status_code=404, detail="账号不存在")
    return account


async def get_video(session: AsyncSession, video_id: str) -> Video:
    video = await session.scalar(select(Video).where(Video.id == video_id))
    if not video:
        raise HTTPException(status_code=404, detail="视频不存在")
    return video


async def seed_database() -> None:
    settings = get_settings()
    async with SessionLocal() as session:
        workspace = await session.scalar(select(Workspace).limit(1))
        if not workspace:
            workspace = Workspace(name="VidInsight 内部工作区")
            session.add(workspace)
            await session.flush()
        admin = await session.scalar(select(User).where(User.email == settings.admin_email))
        if not admin:
            session.add(User(email=settings.admin_email, display_name="系统管理员", password_hash=hash_password(settings.admin_password), role=Role.ADMIN.value))
        existing = set((await session.scalars(select(MetricRequirement.metric_code).where(MetricRequirement.workspace_id == workspace.id))).all())
        for code, scope in REQUIRED_METRICS.items():
            if code not in existing:
                session.add(MetricRequirement(workspace_id=workspace.id, metric_code=code, scope=scope, coverage_status="pending"))
        await session.commit()


@asynccontextmanager
async def lifespan(_: FastAPI):
    await create_schema()
    await seed_database()
    yield


app = FastAPI(title="VidInsight API", version="1.0.0", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=get_settings().cors_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/health/live")
async def health_live() -> dict:
    return {"status": "ok"}


@app.get("/health/ready")
async def health_ready(session: AsyncSession = Depends(get_db)) -> dict:
    await session.execute(select(Workspace.id).limit(1))
    return {"status": "ready"}


@app.post("/api/v1/auth/login")
async def login(body: LoginRequest, session: AsyncSession = Depends(get_db)) -> dict:
    user = await session.scalar(select(User).where(User.email == body.email, User.active.is_(True)))
    if not user or not verify_password(body.password, user.password_hash):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="邮箱或密码错误")
    workspace = await get_workspace(session)
    return {"access_token": create_access_token(user), "token_type": "bearer", "user": {"id": user.id, "email": user.email, "role": user.role}, "workspace_id": workspace.id}


@app.get("/api/v1/auth/me")
async def me(user: User = Depends(current_user), session: AsyncSession = Depends(get_db)) -> dict:
    workspace = await get_workspace(session)
    return {"id": user.id, "email": user.email, "display_name": user.display_name, "role": user.role, "workspace_id": workspace.id}


@app.get("/api/v1/integrations/douyin/preflight")
async def douyin_preflight(_: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    settings = get_settings()
    required = ["DOUYIN_CLIENT_KEY", "DOUYIN_CLIENT_SECRET", "DOUYIN_REDIRECT_URI"]
    configured = {
        "DOUYIN_CLIENT_KEY": bool(settings.douyin_client_key),
        "DOUYIN_CLIENT_SECRET": bool(settings.douyin_client_secret),
        "DOUYIN_REDIRECT_URI": bool(settings.douyin_redirect_uri),
    }
    return {
        "ready_for_oauth": all(configured.values()),
        "configured": configured,
        "requested_scopes": settings.douyin_scope_list,
        "missing": [name for name in required if not configured[name]],
        "field_gaps": ["completion_rate", "video_follow_conversion"],
        "notes": ["Official video query supports play/like/comment/share/forward only with an approved capability and known item_id.", "Do not enable production collection until the approved scope, use case, quota, and field samples have been recorded."],
    }


@app.post("/api/v1/accounts/{account_id}/authorize")
async def authorize_douyin_account(account_id: str, session: AsyncSession = Depends(get_db), _: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    account = await get_account(session, account_id)
    state = new_state()
    session.add(OAuthState(account_id=account.id, state_hash=state_hash(state), expires_at=utcnow().replace(microsecond=0) + timedelta(minutes=10)))
    await session.commit()
    try:
        return {"authorization_url": authorization_url(state), "expires_in": 600}
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc


@app.get("/api/v1/integrations/douyin/callback")
async def douyin_oauth_callback(code: str, state: str, scopes: str = "", session: AsyncSession = Depends(get_db)):
    record = await session.scalar(select(OAuthState).where(OAuthState.state_hash == state_hash(state), OAuthState.consumed_at.is_(None), OAuthState.expires_at > utcnow()))
    if not record:
        raise HTTPException(status_code=400, detail="Invalid or expired OAuth state")
    account = await get_account(session, record.account_id)
    try:
        payload = await exchange_code(code)
    except (RuntimeError, ValueError, httpx.HTTPError) as exc:
        account.authorization_status = AccountStatus.EXPIRED.value
        account.last_sync_error = f"OAuth exchange failed: {exc.__class__.__name__}"
        await session.commit()
        raise HTTPException(status_code=502, detail="Douyin OAuth token exchange failed") from exc
    if payload["open_id"] != account.platform_open_id:
        raise HTTPException(status_code=409, detail="Authorized Douyin account does not match the selected account")
    account.access_token_encrypted = encrypt(payload["access_token"])
    account.refresh_token_encrypted = encrypt(payload["refresh_token"])
    account.token_expires_at = expiration(payload.get("expires_in"))
    account.refresh_token_expires_at = expiration(payload.get("refresh_expires_in"))
    account.authorized_scopes = scopes or payload.get("scope", "")
    account.authorization_status = AccountStatus.ACTIVE.value
    account.last_sync_error = None
    record.consumed_at = utcnow()
    await session.commit()
    settings = get_settings()
    return RedirectResponse(f"{settings.frontend_app_url.rstrip('/')}/accounts/{account.id}?authorized=1", status_code=303)


@app.get("/api/v1/accounts")
async def list_accounts(session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> list[dict]:
    accounts = (await session.scalars(select(Account).order_by(Account.created_at.desc()))).all()
    return [serialize_account(account) for account in accounts]


@app.post("/api/v1/accounts", status_code=status.HTTP_201_CREATED)
async def create_account(body: AccountCreate, session: AsyncSession = Depends(get_db), user: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    workspace = await get_workspace(session)
    existing = await session.scalar(select(Account).where(Account.workspace_id == workspace.id, Account.platform_open_id == body.platform_open_id))
    if existing:
        raise HTTPException(status_code=409, detail="账号已存在")
    account = Account(workspace_id=workspace.id, platform_open_id=body.platform_open_id, nickname=body.nickname, alias=body.alias, priority=body.priority)
    session.add(account)
    session.add(AuditLog(workspace_id=workspace.id, user_id=user.id, action="account.create", resource_type="account", resource_id=account.id))
    await session.commit()
    await session.refresh(account)
    return serialize_account(account)


@app.post("/api/v1/accounts/batch")
async def batch_accounts(items: list[AccountCreate], session: AsyncSession = Depends(get_db), user: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    if not items or len(items) > 100:
        raise HTTPException(status_code=422, detail="批量录入数量需在 1 到 100 之间")
    workspace = await get_workspace(session)
    results = []
    for index, item in enumerate(items, start=1):
        exists = await session.scalar(select(Account.id).where(Account.workspace_id == workspace.id, Account.platform_open_id == item.platform_open_id))
        if exists:
            results.append({"row_number": index, "status": "failed", "error": "账号已存在"})
            continue
        account = Account(workspace_id=workspace.id, platform_open_id=item.platform_open_id, nickname=item.nickname, alias=item.alias, priority=item.priority)
        session.add(account)
        results.append({"row_number": index, "status": "created", "account_id": account.id})
    session.add(AuditLog(workspace_id=workspace.id, user_id=user.id, action="account.batch_create", resource_type="account", details={"count": len(items)}))
    await session.commit()
    return {"items": results}


@app.get("/api/v1/accounts/{account_id}")
async def account_detail(account_id: str, session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> dict:
    return serialize_account(await get_account(session, account_id))


@app.patch("/api/v1/accounts/{account_id}")
async def patch_account(account_id: str, body: AccountPatch, session: AsyncSession = Depends(get_db), _: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    account = await get_account(session, account_id)
    for name, value in body.model_dump(exclude_none=True).items():
        setattr(account, name, value)
    await session.commit()
    return serialize_account(account)


@app.post("/api/v1/accounts/{account_id}/enable")
async def enable_account(account_id: str, session: AsyncSession = Depends(get_db), _: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    account = await get_account(session, account_id)
    account.collection_enabled = True
    account.authorization_status = AccountStatus.ACTIVE.value
    await session.commit()
    return serialize_account(account)


@app.post("/api/v1/accounts/{account_id}/disable")
async def disable_account(account_id: str, session: AsyncSession = Depends(get_db), _: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    account = await get_account(session, account_id)
    account.collection_enabled = False
    account.authorization_status = AccountStatus.DISABLED.value
    await session.commit()
    return serialize_account(account)


@app.delete("/api/v1/accounts/{account_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_account(account_id: str, session: AsyncSession = Depends(get_db), _: User = Depends(require_role(Role.ADMIN.value))):
    account = await get_account(session, account_id)
    await session.delete(account)
    await session.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@app.post("/api/v1/accounts/{account_id}/sync", status_code=status.HTTP_202_ACCEPTED)
async def sync_account(account_id: str, session: AsyncSession = Depends(get_db), _: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    account = await get_account(session, account_id)
    if not account.collection_enabled:
        raise HTTPException(status_code=409, detail="账号已停用")
    job = SyncJob(workspace_id=account.workspace_id, account_id=account.id, trigger="manual", status=JobStatus.QUEUED.value)
    session.add(job)
    await session.commit()
    try:
        from app.tasks.celery_app import collect_account
        collect_account.delay(job.id)
    except Exception as exc:
        job.error_message = f"Task is queued; retry when Redis is available ({exc.__class__.__name__})."
        await session.commit()
    return {"job_id": job.id, "status": job.status}


@app.get("/api/v1/sync-jobs/{job_id}")
async def sync_job_detail(job_id: str, session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> dict:
    job = await session.scalar(select(SyncJob).where(SyncJob.id == job_id))
    if not job:
        raise HTTPException(status_code=404, detail="Sync job not found")
    return {"id": job.id, "account_id": job.account_id, "trigger": job.trigger, "status": job.status, "records_written": job.records_written, "error_message": job.error_message, "created_at": job.created_at, "completed_at": job.completed_at}


@app.get("/api/v1/sync-jobs")
async def list_sync_jobs(account_id: str | None = None, limit: int = Query(default=30, ge=1, le=100), session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> list[dict]:
    query = select(SyncJob).order_by(SyncJob.created_at.desc()).limit(limit)
    if account_id:
        query = query.where(SyncJob.account_id == account_id)
    rows = (await session.scalars(query)).all()
    return [{"id": job.id, "account_id": job.account_id, "trigger": job.trigger, "status": job.status, "records_written": job.records_written, "error_message": job.error_message, "created_at": job.created_at, "completed_at": job.completed_at} for job in rows]


@app.post("/api/v1/accounts/{account_id}/snapshots", status_code=status.HTTP_201_CREATED)
async def import_account_snapshot(account_id: str, body: AccountSnapshotImport, session: AsyncSession = Depends(get_db), _: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    account = await get_account(session, account_id)
    snapshot = AccountMetricSnapshot(workspace_id=account.workspace_id, account_id=account.id, follower_count=body.follower_count, collected_at=body.collected_at, source=body.source)
    session.add(snapshot)
    await session.commit()
    return {"id": snapshot.id, "status": "created"}


@app.post("/api/v1/accounts/{account_id}/videos/import", status_code=status.HTTP_201_CREATED)
async def import_video(account_id: str, body: VideoImport, session: AsyncSession = Depends(get_db), _: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    account = await get_account(session, account_id)
    video = await session.scalar(select(Video).where(Video.account_id == account.id, Video.platform_item_id == body.platform_item_id))
    if not video:
        video = Video(workspace_id=account.workspace_id, account_id=account.id, platform_item_id=body.platform_item_id, title=body.title, published_at=body.published_at, duration_ms=body.duration_ms)
        session.add(video)
        await session.flush()
    else:
        video.title = body.title
        video.published_at = body.published_at
        video.duration_ms = body.duration_ms
    metric = VideoMetricSnapshot(workspace_id=account.workspace_id, video_id=video.id, play_count=body.play_count, like_count=body.like_count, comment_count=body.comment_count, share_count=body.share_count, forward_count=body.forward_count, completion_rate=body.completion_rate, completion_sample_count=body.completion_sample_count, video_follow_conversion=body.video_follow_conversion, collected_at=body.collected_at, source=body.source)
    session.add(metric)
    account.last_synced_at = body.collected_at
    await session.commit()
    return {"video_id": video.id, "metric_id": metric.id}


@app.post("/api/v1/accounts/{account_id}/data/import", status_code=status.HTTP_201_CREATED)
async def import_account_data(account_id: str, body: AccountDataImport, session: AsyncSession = Depends(get_db), _: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    """Append a validated batch from an approved export or a manual backfill.

    The transaction is all-or-nothing for invalid input; duplicate snapshot points
    are reported and preserved rather than overwriting historical data.
    """
    account = await get_account(session, account_id)
    created = {"account_snapshots": 0, "video_metrics": 0}
    skipped = {"account_snapshots": 0, "video_metrics": 0}
    if body.account_snapshot:
        item = body.account_snapshot
        duplicate = await session.scalar(select(AccountMetricSnapshot.id).where(AccountMetricSnapshot.account_id == account.id, AccountMetricSnapshot.collected_at == item.collected_at, AccountMetricSnapshot.source == item.source))
        if duplicate:
            skipped["account_snapshots"] += 1
        else:
            session.add(AccountMetricSnapshot(workspace_id=account.workspace_id, account_id=account.id, follower_count=item.follower_count, collected_at=item.collected_at, source=item.source))
            created["account_snapshots"] += 1
    for item in body.videos:
        video = await session.scalar(select(Video).where(Video.account_id == account.id, Video.platform_item_id == item.platform_item_id))
        if not video:
            video = Video(workspace_id=account.workspace_id, account_id=account.id, platform_item_id=item.platform_item_id, title=item.title, published_at=item.published_at, duration_ms=item.duration_ms)
            session.add(video)
            await session.flush()
        duplicate = await session.scalar(select(VideoMetricSnapshot.id).where(VideoMetricSnapshot.video_id == video.id, VideoMetricSnapshot.collected_at == item.collected_at, VideoMetricSnapshot.source == item.source))
        if duplicate:
            skipped["video_metrics"] += 1
            continue
        session.add(VideoMetricSnapshot(workspace_id=account.workspace_id, video_id=video.id, play_count=item.play_count, like_count=item.like_count, comment_count=item.comment_count, share_count=item.share_count, forward_count=item.forward_count, completion_rate=item.completion_rate, completion_sample_count=item.completion_sample_count, video_follow_conversion=item.video_follow_conversion, collected_at=item.collected_at, source=item.source))
        account.last_synced_at = item.collected_at
        created["video_metrics"] += 1
    await session.commit()
    return {"created": created, "skipped": skipped}


@app.get("/api/v1/dashboard/overview")
async def overview(start_at: datetime, end_at: datetime, account_ids: list[str] = Query(default=[]), session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> dict:
    data = await period_metrics(session, start_at, end_at, account_ids)
    return {"range": {"start_at": start_at, "end_at": end_at}, **data}


@app.get("/api/v1/dashboard/trends")
async def trends(start_at: datetime, end_at: datetime, account_ids: list[str] = Query(default=[]), session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> list[dict]:
    metrics = await session.scalars(select(VideoMetricSnapshot).where(VideoMetricSnapshot.collected_at >= start_at, VideoMetricSnapshot.collected_at < end_at).order_by(VideoMetricSnapshot.collected_at))
    points: dict[str, dict] = {}
    for item in metrics:
        video = await get_video(session, item.video_id)
        if account_ids and video.account_id not in account_ids:
            continue
        key = item.collected_at.date().isoformat()
        point = points.setdefault(key, {"date": key, "plays": 0, "likes": 0, "comments": 0, "forwards": 0})
        point["plays"] += item.play_count
        point["likes"] += item.like_count
        point["comments"] += item.comment_count
        point["forwards"] += item.forward_count
    return list(points.values())


@app.get("/api/v1/accounts/{account_id}/analytics")
async def account_analytics(account_id: str, start_at: datetime, end_at: datetime, session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> dict:
    account = await get_account(session, account_id)
    account_metrics = await period_metrics(session, start_at, end_at, [account.id])
    account_trends = await trends(start_at, end_at, [account.id], session)
    videos = await list_videos(account_id=account.id, session=session, _=_)
    diagnoses = (await session.scalars(select(Diagnosis).join(Video, Diagnosis.video_id == Video.id).where(Video.account_id == account.id).order_by(Diagnosis.created_at.desc()))).all()
    by_video: dict[str, list[dict]] = {}
    for item in diagnoses:
        by_video.setdefault(item.video_id, []).append({"rule_code": item.rule_code, "title": item.title, "suggestion": item.suggestion, "created_at": item.created_at})
    for video in videos:
        video["diagnoses"] = by_video.get(video["id"], [])
    return {"account": serialize_account(account), "range": {"start_at": start_at, "end_at": end_at}, **account_metrics, "trends": account_trends, "videos": videos}


@app.get("/api/v1/videos")
async def list_videos(account_id: str | None = None, q: str | None = None, start_at: datetime | None = None, end_at: datetime | None = None, sort: str = "-published_at", session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> list[dict]:
    query = select(Video)
    if account_id:
        query = query.where(Video.account_id == account_id)
    if q:
        query = query.where(Video.title.ilike(f"%{q}%"))
    if start_at:
        query = query.where(Video.published_at >= start_at)
    if end_at:
        query = query.where(Video.published_at < end_at)
    videos = (await session.scalars(query.order_by(Video.published_at.desc()))).all()
    values = [await serialize_video(session, video) for video in videos]
    key = {"play_count": "play_count", "like_count": "like_count", "comment_count": "comment_count", "forward_count": "forward_count"}.get(sort.lstrip("-"))
    if key:
        values.sort(key=lambda item: item["latest_metrics"].get(key, 0), reverse=sort.startswith("-"))
    return values


@app.get("/api/v1/videos/{video_id}")
async def video_detail(video_id: str, session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> dict:
    return await serialize_video(session, await get_video(session, video_id), include_history=True)


@app.post("/api/v1/videos/{video_id}/diagnose", status_code=status.HTTP_202_ACCEPTED)
async def diagnose_video(video_id: str, session: AsyncSession = Depends(get_db), _: User = Depends(require_role(Role.ADMIN.value, Role.OPERATOR.value))) -> dict:
    video = await get_video(session, video_id)
    metric = await latest_metric(session, video.id)
    if not metric or metric.completion_rate is None:
        raise HTTPException(status_code=422, detail="该视频缺少必需的指标数据")
    all_metrics = (await session.scalars(select(VideoMetricSnapshot).where(VideoMetricSnapshot.workspace_id == video.workspace_id))).all()
    baseline = calculate_baseline(all_metrics)
    metric_data = metric_dict(metric)
    await session.execute(Diagnosis.__table__.delete().where(Diagnosis.video_id == video.id))
    created = []
    for rule, evidence in build_diagnoses(metric_data, baseline):
        item = Diagnosis(workspace_id=video.workspace_id, video_id=video.id, rule_code=rule.code, title=rule.title, suggestion=rule.suggestion, evidence=evidence)
        session.add(item)
        created.append(rule.code)
    await session.commit()
    return {"video_id": video.id, "rules": created}


@app.get("/api/v1/diagnoses")
async def list_diagnoses(account_id: str | None = None, session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> list[dict]:
    rows = (await session.scalars(select(Diagnosis).order_by(Diagnosis.created_at.desc()))).all()
    result = []
    for item in rows:
        video = await get_video(session, item.video_id)
        if account_id and video.account_id != account_id:
            continue
        result.append({"id": item.id, "video_id": item.video_id, "rule_code": item.rule_code, "title": item.title, "suggestion": item.suggestion, "evidence": item.evidence, "created_at": item.created_at})
    return result


@app.get("/api/v1/metrics/coverage")
async def metrics_coverage(session: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> dict:
    workspace = await get_workspace(session)
    videos = (await session.scalars(select(Video).where(Video.workspace_id == workspace.id))).all()
    accounts = (await session.scalars(select(Account).where(Account.workspace_id == workspace.id))).all()
    report = []
    for code, scope in REQUIRED_METRICS.items():
        if scope == "video":
            total = len(videos)
            covered = 0
            for video in videos:
                metric = await latest_metric(session, video.id)
                if metric and getattr(metric, code, None) is not None and metric.source != "model_estimate":
                    covered += 1
        else:
            total = len(accounts)
            covered = 0
            for account in accounts:
                snapshots = (await session.scalars(select(AccountMetricSnapshot).where(AccountMetricSnapshot.account_id == account.id).order_by(AccountMetricSnapshot.collected_at.desc()).limit(2))).all()
                if len(snapshots) >= 2:
                    covered += 1
        report.append({"metric_code": code, "scope": scope, "covered": covered, "total": total, "coverage_rate": covered / total if total else 0})
    ready = bool(report) and all(item["total"] > 0 and item["coverage_rate"] == 1 for item in report)
    return {"ready_for_acceptance": ready, "metrics": report}


@app.post("/api/v1/exports", status_code=status.HTTP_201_CREATED)
async def create_export(body: ExportRequest, session: AsyncSession = Depends(get_db), user: User = Depends(current_user)) -> dict:
    workspace = await get_workspace(session)
    job = ExportJob(workspace_id=workspace.id, requested_by=user.id, filters=body.model_dump(mode="json"), status=JobStatus.RUNNING.value)
    session.add(job)
    await session.flush()
    data = await period_metrics(session, body.start_at, body.end_at, body.account_ids)
    rows = await export_video_rows(session, body.start_at, body.end_at, body.account_ids)
    path = write_xlsx(job.id, data, rows)
    job.status = JobStatus.SUCCEEDED.value
    job.file_name = path.name
    job.file_path = str(path)
    await session.commit()
    return {"id": job.id, "status": job.status, "file_name": job.file_name}


@app.get("/api/v1/exports/{job_id}/download")
async def download_export(job_id: str, session: AsyncSession = Depends(get_db), _: User = Depends(current_user)):
    job = await session.scalar(select(ExportJob).where(ExportJob.id == job_id, ExportJob.status == JobStatus.SUCCEEDED.value))
    if not job or not job.file_path or not Path(job.file_path).is_file():
        raise HTTPException(status_code=404, detail="导出文件不存在或已过期")
    return FileResponse(job.file_path, filename=job.file_name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def serialize_account(account: Account) -> dict:
    return {"id": account.id, "platform_open_id": account.platform_open_id, "nickname": account.nickname, "alias": account.alias, "authorization_status": account.authorization_status, "collection_enabled": account.collection_enabled, "priority": account.priority, "last_synced_at": account.last_synced_at}


async def latest_metric(session: AsyncSession, video_id: str) -> VideoMetricSnapshot | None:
    return await session.scalar(select(VideoMetricSnapshot).where(VideoMetricSnapshot.video_id == video_id).order_by(VideoMetricSnapshot.collected_at.desc()).limit(1))


def metric_dict(metric: VideoMetricSnapshot) -> dict:
    return {"play_count": metric.play_count, "like_count": metric.like_count, "comment_count": metric.comment_count, "share_count": metric.share_count, "forward_count": metric.forward_count, "completion_rate": float(metric.completion_rate or 0), "video_follow_conversion": float(metric.video_follow_conversion) if metric.video_follow_conversion is not None else None}


async def serialize_video(session: AsyncSession, video: Video, include_history: bool = False) -> dict:
    metric = await latest_metric(session, video.id)
    result = {"id": video.id, "account_id": video.account_id, "platform_item_id": video.platform_item_id, "title": video.title, "published_at": video.published_at, "duration_ms": video.duration_ms, "latest_metrics": metric_dict(metric) if metric else {}}
    if include_history:
        history = (await session.scalars(select(VideoMetricSnapshot).where(VideoMetricSnapshot.video_id == video.id).order_by(VideoMetricSnapshot.collected_at))).all()
        result["history"] = [{**metric_dict(item), "collected_at": item.collected_at, "source": item.source} for item in history]
    return result


def calculate_baseline(metrics: list[VideoMetricSnapshot]) -> dict[str, float]:
    plays = sorted([item.play_count for item in metrics] or [0])
    completions = sorted([float(item.completion_rate) for item in metrics if item.completion_rate is not None] or [0])
    interactions = sorted([(item.like_count + item.comment_count + item.share_count) / item.play_count for item in metrics if item.play_count] or [0])
    follows = sorted([float(item.video_follow_conversion) for item in metrics if item.video_follow_conversion is not None] or [0])
    def p(values: list[float], fraction: float) -> float:
        return values[min(len(values) - 1, int((len(values) - 1) * fraction))]
    return {"play_p30": p(plays, .3), "play_p70": p(plays, .7), "completion_p30": p(completions, .3), "completion_p70": p(completions, .7), "interaction_p30": p(interactions, .3), "follow_p30": p(follows, .3)}


async def period_metrics(session: AsyncSession, start_at: datetime, end_at: datetime, account_ids: list[str]) -> dict:
    videos = (await session.scalars(select(Video))).all()
    selected = [video for video in videos if not account_ids or video.account_id in account_ids]
    totals = {"play_delta": 0, "like_delta": 0, "comment_delta": 0, "forward_delta": 0}
    completion_num = completion_den = 0.0
    for video in selected:
        snapshots = (await session.scalars(select(VideoMetricSnapshot).where(VideoMetricSnapshot.video_id == video.id).order_by(VideoMetricSnapshot.collected_at))).all()
        before = [row for row in snapshots if row.collected_at < start_at]
        inside = [row for row in snapshots if start_at <= row.collected_at < end_at]
        if not inside:
            continue
        first = before[-1] if before else inside[0]
        last = inside[-1]
        totals["play_delta"] += max(0, last.play_count - first.play_count)
        totals["like_delta"] += max(0, last.like_count - first.like_count)
        totals["comment_delta"] += max(0, last.comment_count - first.comment_count)
        totals["forward_delta"] += max(0, last.forward_count - first.forward_count)
        if last.completion_rate is not None:
            weight = last.completion_sample_count or 1
            completion_num += float(last.completion_rate) * weight
            completion_den += weight
    accounts = (await session.scalars(select(Account))).all()
    follower_delta = 0
    for account in accounts:
        if account_ids and account.id not in account_ids:
            continue
        snapshots = (await session.scalars(select(AccountMetricSnapshot).where(AccountMetricSnapshot.account_id == account.id).order_by(AccountMetricSnapshot.collected_at))).all()
        before = [row for row in snapshots if row.collected_at < start_at]
        inside = [row for row in snapshots if start_at <= row.collected_at < end_at]
        if inside:
            follower_delta += max(0, inside[-1].follower_count - (before[-1].follower_count if before else inside[0].follower_count))
    return {"metrics": {**totals, "follower_delta": follower_delta, "avg_completion_rate": completion_num / completion_den if completion_den else None}, "account_count": len(accounts), "video_count": len(selected)}


async def export_video_rows(session: AsyncSession, start_at: datetime, end_at: datetime, account_ids: list[str]) -> list[dict]:
    rows = []
    for video in (await session.scalars(select(Video))).all():
        if account_ids and video.account_id not in account_ids:
            continue
        snapshots = (await session.scalars(select(VideoMetricSnapshot).where(VideoMetricSnapshot.video_id == video.id, VideoMetricSnapshot.collected_at >= start_at, VideoMetricSnapshot.collected_at < end_at).order_by(VideoMetricSnapshot.collected_at.desc()).limit(1))).all()
        if snapshots:
            rows.append({"id": video.id, "account_id": video.account_id, "platform_item_id": video.platform_item_id, "title": video.title, "published_at": video.published_at, "latest_metrics": metric_dict(snapshots[0])})
    return rows


def write_xlsx(job_id: str, summary: dict, rows: list[dict]) -> Path:
    settings = get_settings()
    path = settings.export_dir / f"vidinsight-{job_id}.xlsx"
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总"
    summary_sheet.append(["指标", "数值"])
    for key, value in summary["metrics"].items():
        summary_sheet.append([key, value])
    detail = workbook.create_sheet("视频明细")
    detail.append(["视频 ID", "标题", "账号 ID", "发布时间", "播放", "点赞", "评论", "转发", "完播率", "来源"])
    for row in rows:
        metric = row["latest_metrics"]
        detail.append([row["platform_item_id"], row["title"], row["account_id"], row["published_at"].isoformat(), metric.get("play_count"), metric.get("like_count"), metric.get("comment_count"), metric.get("forward_count"), metric.get("completion_rate"), "latest_snapshot"])
    glossary = workbook.create_sheet("口径说明")
    glossary.append(["指标", "口径"])
    glossary.append(["周期增量", "周期结束最后快照减去周期开始前最后快照"])
    glossary.append(["平均完播率", "按 completion_sample_count 加权；缺失样本数时简单平均"])
    workbook.save(path)
    return path
